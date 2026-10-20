"""Text extraction and similarity scoring; scores are not proof of plagiarism."""
import io
import re
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


def normalize_text(text):
    return re.sub(r'\s+', ' ', text.lower()).strip()


def extract_text_from_file(file_obj, suffix):
    raw = file_obj.read()
    if suffix == '.txt':
        return raw.decode('utf-8-sig', errors='replace')
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        return '\n'.join(page.extract_text() or '' for page in reader.pages)
    if suffix == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(' | '.join(cell.text for cell in row.cells))
        return '\n'.join(parts)
    raise ValueError('Unsupported file type')


def _score(vectorizer, a, b):
    try:
        matrix = vectorizer.fit_transform([a, b])
    except ValueError:
        # Very short or punctuation-only text may produce no valid n-grams.
        return 0.0
    return round(float(cosine_similarity(matrix[0:1], matrix[1:2])[0, 0]) * 100, 2)


def score_pair(query, reference):
    query, reference = normalize_text(query), normalize_text(reference)
    if not query or not reference:
        return {'word_score': 0.0, 'phrase_score': 0.0, 'combined_score': 0.0}
    word = _score(TfidfVectorizer(analyzer='word', ngram_range=(1, 2), sublinear_tf=True,
                                  token_pattern=r'(?u)\b\w+\b', min_df=1), query, reference)
    char = _score(TfidfVectorizer(analyzer='char_wb', ngram_range=(3, 5), sublinear_tf=True,
                                  min_df=1), query, reference)
    return {'word_score': word, 'phrase_score': char,
            'combined_score': round(0.7 * word + 0.3 * char, 2)}


def compare_documents(query, references):
    if not query.strip():
        raise ValueError('Query text cannot be empty.')
    results = []
    for ref in references:
        text = ref.get('text', '')
        if not isinstance(text, str) or not text.strip():
            continue
        results.append({'name': ref.get('name', 'Reference'), **score_pair(query, text),
                        'query_words': len(normalize_text(query).split()),
                        'reference_words': len(normalize_text(text).split())})
    if not results:
        raise ValueError('No readable reference text was found.')
    return sorted(results, key=lambda item: item['combined_score'], reverse=True)
